import os
import sys
import json
import re
import time
import math
from urllib.parse import quote, unquote
import requests
import openpyxl

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
JSON_PATH = os.path.join(PROJECT_ROOT, "data", "processed", "debi_dorshon.json")
EXCEL_PATH = os.path.join(PROJECT_ROOT, "data", "raw", "Debi-Dorshon.xlsx")
DEMO_JSON_PATH = os.path.join(PROJECT_ROOT, "data", "processed", "debi_dorshon_demo.json")

session = requests.Session()
session.headers.update({
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
    'Accept-Language': 'en-US,en;q=0.9'
})

def is_valid_kolkata(lat, lng):
    if lat is None or lng is None:
        return False
    return (22.35 <= lat <= 22.80 and 88.20 <= lng <= 88.60)

def parse_gmaps_response(text):
    if text.startswith(")]}'"):
        text = text[4:].strip()
    try:
        data = json.loads(text)
    except Exception:
        return None, None, "json_error"

    # Method 1: Check places list data[0][1]
    if len(data) > 0 and isinstance(data[0], list) and len(data[0]) > 1 and isinstance(data[0][1], list):
        for item in data[0][1]:
            if isinstance(item, list) and len(item) > 14 and isinstance(item[14], list):
                info = item[14]
                if len(info) > 9 and isinstance(info[9], list) and len(info[9]) >= 4:
                    lat, lng = info[9][2], info[9][3]
                    if isinstance(lat, (int, float)) and isinstance(lng, (int, float)):
                        if is_valid_kolkata(lat, lng):
                            return float(lat), float(lng), "gmaps_place_pin"
                if len(info) > 208 and isinstance(info[208], list) and len(info[208]) > 0 and isinstance(info[208][0], list) and len(info[208][0]) >= 4:
                    lat, lng = info[208][0][2], info[208][0][3]
                    if isinstance(lat, (int, float)) and isinstance(lng, (int, float)):
                        if is_valid_kolkata(lat, lng):
                            return float(lat), float(lng), "gmaps_place_pin_208"

    # Method 2: Check viewport center data[1][0]
    if len(data) > 1 and isinstance(data[1], list) and len(data[1]) > 0 and isinstance(data[1][0], list):
        v = data[1][0]
        if len(v) >= 3:
            lng, lat = v[1], v[2]
            if isinstance(lat, (int, float)) and isinstance(lng, (int, float)):
                if is_valid_kolkata(lat, lng):
                    return float(lat), float(lng), "gmaps_viewport_center"

    return None, None, "not_found"

def query_google_maps(query):
    try:
        url = f"https://www.google.com/search?tbm=map&authuser=0&hl=en&gl=in&q={quote(query)}"
        r = session.get(url, timeout=10)
        return parse_gmaps_response(r.text)
    except Exception as e:
        return None, None, f"error: {e}"

def resolve_gmaps_url(url):
    if not url:
        return None, None, None, None
    try:
        r = session.get(url, allow_redirects=True, timeout=10)
        final = unquote(r.url)
        m = re.search(r'!(?:8m2!)?3d(-?\d+(?:\.\d+)?)!4d(-?\d+(?:\.\d+)?)', final)
        if m:
            lat, lng = float(m.group(1)), float(m.group(2))
            if is_valid_kolkata(lat, lng):
                return lat, lng, "gmaps_link_pin", final
        m = re.search(r'@(-?\d+(?:\.\d+)?),(-?\d+(?:\.\d+)?)', final)
        if m:
            lat, lng = float(m.group(1)), float(m.group(2))
            if is_valid_kolkata(lat, lng):
                return lat, lng, "gmaps_link_center", final
        m_place = re.search(r'/place/([^/@?]+)', final)
        if m_place:
            place_name = m_place.group(1).replace('+', ' ')
            return None, None, "place_address", place_name
    except Exception as e:
        return None, None, f"error: {e}", None
    return None, None, "unresolved", final

def load_excel_links():
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    excel_pandals = []
    for sheetname in wb.sheetnames:
        sheet = wb[sheetname]
        for r in range(1, sheet.max_row+1):
            for c in range(1, sheet.max_column+1):
                val = str(sheet.cell(r, c).value or '').strip()
                if val in ['Sl.no.', 'Sl. No.', 'Sl.No.']:
                    dr = r + 1
                    while dr <= sheet.max_row:
                        name = sheet.cell(dr, c+1).value
                        if not name or str(sheet.cell(dr, c).value).strip() in ['Sl.no.', 'Sl. No.']:
                            break
                        loc_cell = sheet.cell(dr, c+2)
                        link = loc_cell.hyperlink.target if loc_cell.hyperlink else None
                        excel_pandals.append({
                            'sheet': sheetname,
                            'name': str(name).strip(),
                            'link': link
                        })
                        dr += 1
    return excel_pandals

def main():
    with open(JSON_PATH, "r", encoding="utf-8") as f:
        original_pandals = json.load(f)

    excel_pandals = load_excel_links()
    print(f"Loaded {len(original_pandals)} pandals from debi_dorshon.json")
    print(f"Loaded {len(excel_pandals)} entries from Excel")

    demo_pandals = []
    resolved_count = 0
    unresolved_list = []

    for idx, p in enumerate(original_pandals):
        name = p["name"]
        region = p.get("region", "")
        cluster = p.get("cluster", "")

        # 1. Match with Excel Google Maps link
        link = None
        clean_name = re.sub(r'\(.*?\)', '', name).strip()
        for ep in excel_pandals:
            ep_clean = re.sub(r'\(.*?\)', '', ep["name"]).strip()
            if name.lower() == ep["name"].lower() or clean_name.lower() == ep_clean.lower() or name.lower() in ep["name"].lower() or ep["name"].lower() in name.lower():
                if ep["link"]:
                    link = ep["link"]
                    break

        lat, lng, method = None, None, None

        if link:
            l_lat, l_lng, l_method, l_info = resolve_gmaps_url(link)
            if l_lat and l_lng:
                lat, lng, method = l_lat, l_lng, f"gmaps_link ({l_method})"
            elif l_method == "place_address" and l_info:
                a_lat, a_lng, a_method = query_google_maps(l_info)
                if a_lat and a_lng:
                    lat, lng, method = a_lat, a_lng, f"gmaps_link_address ({a_method})"

        # 2. Search on Google Maps if not found
        if not lat or not lng:
            search_candidates = [
                f"{name} Durga Puja Pandal Kolkata",
                f"{name} Durga Puja {cluster} Kolkata",
                f"{name} Puja Pandal Kolkata",
                f"{clean_name} Durga Puja Kolkata",
                f"{name} {cluster} Kolkata",
                f"{name} Kolkata"
            ]
            for query in search_candidates:
                s_lat, s_lng, s_method = query_google_maps(query)
                if s_lat and s_lng:
                    lat, lng, method = s_lat, s_lng, f"gmaps_search ({query})"
                    break
                time.sleep(0.08)

        if lat and lng:
            resolved_count += 1
            demo_item = {
                "name": name,
                "region": region,
                "cluster": cluster,
                "location": {
                    "latitude": round(lat, 7),
                    "longitude": round(lng, 7)
                },
                "nearest_metro": p.get("nearest_metro"),
                "nearest_stations": p.get("nearest_stations", []),
                "nearest_ferry": p.get("nearest_ferry")
            }
            demo_pandals.append(demo_item)
        else:
            unresolved_list.append(name)
            demo_item = {
                "name": name,
                "region": region,
                "cluster": cluster,
                "location": {
                    "latitude": None,
                    "longitude": None
                },
                "nearest_metro": p.get("nearest_metro"),
                "nearest_stations": p.get("nearest_stations", []),
                "nearest_ferry": p.get("nearest_ferry")
            }
            demo_pandals.append(demo_item)

        time.sleep(0.05)

    print(f"Resolved: {resolved_count}/{len(original_pandals)}")
    with open(DEMO_JSON_PATH, "w", encoding="utf-8") as f:
        json.dump(demo_pandals, f, indent=2, ensure_ascii=False)
    print(f"Saved demo file: {DEMO_JSON_PATH}")

if __name__ == "__main__":
    main()
